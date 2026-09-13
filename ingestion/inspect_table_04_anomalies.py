from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

FILE_PATH = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "pbs"
    / "census_2023"
    / "table_04"
    / "table_4_punjab_districts.xlsx"
)


def print_rows(worksheet, start_row, end_row):
    for row_number in range(start_row, end_row + 1):
        values = [
            worksheet.cell(row=row_number, column=column).value
            for column in range(1, 17)
        ]

        print(
            f"Row {row_number}: "
            f"{tuple(values)}"
        )


def main():
    workbook = load_workbook(
        FILE_PATH,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook["punjab"]

    print("RAJANPUR AREA")
    print("=" * 80)

    print_rows(
        worksheet,
        12928,
        12938,
    )

    print()
    print("COLUMN 16 ANOMALY")
    print("=" * 80)

    print_rows(
        worksheet,
        2983,
        2989,
    )

    workbook.close()


if __name__ == "__main__":
    main()