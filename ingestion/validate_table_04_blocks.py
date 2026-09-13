from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

FILE_PATH = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "pbs"
    / "census_2023"
    / "table_04"
    / "table_4_punjab_districts.xlsx"
)


def is_blank(value):
    return value is None or value == ""


def is_geography_heading(values):
    first_value = values[0]

    if not isinstance(first_value, str):
        return False

    text = first_value.strip().upper()

    if not all(is_blank(value) for value in values[1:13]):
        return False

    return (
        text.endswith(" DISTRICT")
        or text.endswith(" TEHSIL")
        or text == "DE-EXCLUDED AREA RAJANPUR"
    )


def main():
    workbook = load_workbook(
        FILE_PATH,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook["punjab"]

    blocks = []

    rows = list(worksheet.iter_rows(values_only=True))

    for index, row in enumerate(rows):
        values = list(row)

        if not is_geography_heading(values):
            continue

        geography_name = str(values[0]).strip()

        data_rows = rows[index + 1:index + 93]

        labels = [
            str(data_row[0]).strip()
            if data_row[0] is not None
            else None
            for data_row in data_rows
        ]

        blocks.append(
            {
                "geography": geography_name,
                "start_row": index + 1,
                "labels": labels,
                "row_count": len(data_rows),
            }
        )

    reference = blocks[0]["labels"]

    mismatches = []

    for block in blocks:
        if block["labels"] != reference:
            mismatches.append(block)

    print("BLOCK TEMPLATE VALIDATION")
    print("=" * 70)

    print()
    print(f"Geography blocks detected: {len(blocks)}")

    print()
    print(
        f"Rows expected per geography: "
        f"{len(reference)}"
    )

    print()
    print(
        f"Blocks matching first block exactly: "
        f"{len(blocks) - len(mismatches)}"
    )

    print()
    print(
        f"Blocks with different labels/order: "
        f"{len(mismatches)}"
    )

    if mismatches:
        print()
        print("Mismatched blocks:")

        for block in mismatches:
            print(
                f"Row {block['start_row']}: "
                f"{block['geography']}"
            )

    print()
    print("Reference block labels:")
    print("-" * 70)

    for number, label in enumerate(reference, start=1):
        print(f"{number:02d}. {label}")

    workbook.close()


if __name__ == "__main__":
    main()