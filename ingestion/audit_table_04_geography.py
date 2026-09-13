from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

FILE_PATH = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "pbs"
    / "census_2023"
    / "table_04"
    / "table_4_punjab_districts.xlsx"
)


def is_blank(value):
    return value is None or value == ""


def main():
    workbook = load_workbook(
        FILE_PATH,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook["punjab"]

    structural_rows = []
    heading_suffixes = Counter()
    column_16_values = []

    for row_number, row in enumerate(
        worksheet.iter_rows(values_only=True),
        start=1,
    ):
        values = list(row)

        first_value = values[0]

        # Find rows where column 1 contains text,
        # but the twelve population columns are empty.
        if isinstance(first_value, str):
            text = first_value.strip()

            if (
                text
                and all(is_blank(value) for value in values[1:13])
            ):
                structural_rows.append(
                    (row_number, text)
                )

                last_word = text.split()[-1].upper()
                heading_suffixes[last_word] += 1

        # Inspect the unexpected value in column 16.
        if len(values) >= 16 and not is_blank(values[15]):
            column_16_values.append(
                (row_number, values[15])
            )

    workbook.close()

    print("GEOGRAPHY / STRUCTURE AUDIT")
    print("=" * 70)

    print()
    print(
        f"Structural heading rows detected: "
        f"{len(structural_rows)}"
    )

    print()
    print("Heading types by final word:")
    for suffix, count in heading_suffixes.most_common():
        print(f"{suffix}: {count}")

    print()
    print("First 40 structural headings:")
    for row_number, text in structural_rows[:40]:
        print(f"Row {row_number}: {text}")

    print()
    print("Last 20 structural headings:")
    for row_number, text in structural_rows[-20:]:
        print(f"Row {row_number}: {text}")

    print()
    print("Non-empty values found in column 16:")
    if column_16_values:
        for row_number, value in column_16_values:
            print(
                f"Row {row_number}: {repr(value)}"
            )
    else:
        print("None")


if __name__ == "__main__":
    main()