from pathlib import Path
import csv

import duckdb
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

MANIFEST_PATH = PROJECT_ROOT / "metadata" / "source_manifest.csv"
RAW_DIR = PROJECT_ROOT / "data" / "raw" / "pbs" / "census_2023" / "table_04"
WAREHOUSE_PATH = PROJECT_ROOT / "warehouse" / "pakistan_census.duckdb"


# Only load the regions we have explicitly validated so far.
TARGET_SOURCE_IDS = {
    "pbs_2023_t04_punjab_districts",
    "pbs_2023_t04_kp_districts",
    "pbs_2023_t04_sindh_districts",
    "pbs_2023_t04_balochistan_districts",
    "pbs_2023_t04_islamabad",
}


def find_file(filename):
    matches = list(RAW_DIR.rglob(filename))

    if not matches:
        raise FileNotFoundError(f"Could not find {filename}")

    return matches[0]


def classify_geography(heading):
    heading_upper = heading.upper()

    if heading_upper.endswith(" DISTRICT"):
        return "district"

    if (
        heading_upper.endswith(" SUB-DIVISION")
        or heading_upper.startswith("SUB-DIVISION ")
    ):
        return "sub_division"

    if (
        heading_upper.endswith(" SUB-TEHSIL")
        or heading_upper.startswith("SUB-TEHSIL ")
    ):
        return "sub_tehsil"

    if heading_upper.endswith(" TEHSIL"):
        return "tehsil"

    if heading_upper.endswith(" TALUKA"):
        return "taluka"

    if heading_upper.endswith(" PROTECTED AREA"):
        return "protected_area"

    if heading_upper.startswith("DE-EXCLUDED AREA"):
        return "de_excluded_area"

    raise ValueError(f"Unknown geography heading: {heading}")


def clean_cell(value):
    if value is None:
        return None

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


with MANIFEST_PATH.open("r", encoding="utf-8-sig", newline="") as file:
    manifest = list(csv.DictReader(file))


sources = [
    row
    for row in manifest
    if row["source_id"].strip() in TARGET_SOURCE_IDS
]


connection = duckdb.connect(str(WAREHOUSE_PATH))

connection.execute("create schema if not exists raw")

connection.execute("""
    create table if not exists raw.pbs_census_2023_table_04 (
        source_id varchar,
        source_file varchar,
        source_sheet varchar,
        source_row_number integer,

        geography_name_raw varchar,
        geography_level varchar,
        district_name_raw varchar,

        age_label_raw varchar,

        all_all_sexes_raw varchar,
        all_male_raw varchar,
        all_female_raw varchar,
        all_transgender_raw varchar,

        rural_all_sexes_raw varchar,
        rural_male_raw varchar,
        rural_female_raw varchar,
        rural_transgender_raw varchar,

        urban_all_sexes_raw varchar,
        urban_male_raw varchar,
        urban_female_raw varchar,
        urban_transgender_raw varchar
    )
""")


for source in sources:

    source_id = source["source_id"].strip()
    filename = source["local_filename"].strip()
    region = source["region"].strip()

    path = find_file(filename)

    print("\n" + "=" * 80)
    print(f"LOADING: {region}")
    print(f"SOURCE:  {source_id}")
    print(f"FILE:    {filename}")

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True
    )

    # Current validated workbooks each contain one relevant worksheet.
    sheet_name = workbook.sheetnames[0]
    worksheet = workbook[sheet_name]

    rows = list(worksheet.iter_rows(values_only=True))

    records = []

    current_parent = None

    for index in range(len(rows) - 1):

        current_first_cell = rows[index][0]
        next_first_cell = rows[index + 1][0]

        if current_first_cell is None or next_first_cell is None:
            continue

        heading = str(current_first_cell).strip()

        if str(next_first_cell).strip().upper() != "ALL AGES":
            continue

        geography_level = classify_geography(heading)

        # District and Protected Area are district-equivalent parents.
        if geography_level in {"district", "protected_area"}:
            current_parent = heading
            district_name_raw = heading

        elif geography_level in {
            "tehsil",
	        "taluka",
            "sub_division",
            "sub_tehsil",
            "de_excluded_area",
        }:
            if current_parent is None:
                raise ValueError(
                    f"No parent geography found for {heading}"
                )

            district_name_raw = current_parent

        else:
            raise ValueError(
                f"Unhandled geography level: {geography_level}"
            )

        # Every geography heading must be followed by 92 census rows.
        observation_rows = rows[index + 1:index + 93]

        if len(observation_rows) != 92:
            raise ValueError(
                f"{heading} has {len(observation_rows)} observation rows, expected 92"
            )

        for offset, row in enumerate(observation_rows, start=1):

            age_label = clean_cell(row[0])

            values = [
                clean_cell(value)
                for value in row[1:13]
            ]

            if len(values) != 12:
                raise ValueError(
                    f"Unexpected population column count for {heading}"
                )

            source_row_number = index + offset + 1

            records.append(
                (
                    source_id,
                    filename,
                    sheet_name,
                    source_row_number,

                    heading,
                    geography_level,
                    district_name_raw,

                    age_label,

                    *values,
                )
            )

    expected_records = sum(
        1
        for index in range(len(rows) - 1)
        if rows[index][0] is not None
        and rows[index + 1][0] is not None
        and str(rows[index + 1][0]).strip().upper() == "ALL AGES"
    ) * 92

    if len(records) != expected_records:
        raise ValueError(
            f"{region}: created {len(records)} records, expected {expected_records}"
        )

    # Safe rerun: replace this source only.
    connection.execute(
        """
        delete from raw.pbs_census_2023_table_04
        where source_id = ?
        """,
        [source_id],
    )

    connection.executemany(
        """
        insert into raw.pbs_census_2023_table_04
        values (
            ?, ?, ?, ?,
            ?, ?, ?,
            ?,
            ?, ?, ?, ?,
            ?, ?, ?, ?,
            ?, ?, ?, ?
        )
        """,
        records,
    )

    print(f"GEOGRAPHY BLOCKS: {expected_records // 92}")
    print(f"RAW RECORDS:      {len(records)}")

    workbook.close()


connection.close()

print("\nLOAD COMPLETE")