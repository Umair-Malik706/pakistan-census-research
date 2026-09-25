from pathlib import Path
import csv

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

MANIFEST_PATH = PROJECT_ROOT / "metadata" / "source_manifest.csv"
RAW_DIR = PROJECT_ROOT / "data" / "raw" / "pbs" / "census_2023" / "table_04"


def find_file(filename):
    matches = list(RAW_DIR.rglob(filename))

    if not matches:
        return None

    return matches[0]


def find_geography_headings(ws):
    headings = []

    column_a = [
        row[0]
        for row in ws.iter_rows(
            min_col=1,
            max_col=1,
            values_only=True
        )
    ]

    for index in range(len(column_a) - 1):

        current_value = column_a[index]
        next_value = column_a[index + 1]

        if current_value is None or next_value is None:
            continue

        current_text = str(current_value).strip()
        next_text = str(next_value).strip().upper()

        if next_text == "ALL AGES":
            headings.append((index + 1, current_text))

    return headings


with MANIFEST_PATH.open("r", encoding="utf-8-sig", newline="") as file:
    manifest = list(csv.DictReader(file))


for source in manifest:

    year = str(source["census_year"]).strip()
    table_number = str(source["table_number"]).strip().lstrip("0")

    if year != "2023" or table_number != "4":
        continue

    region = source["region"].strip()


    filename = source["local_filename"].strip()
    source_id = source["source_id"].strip()

    path = find_file(filename)

    print("\n" + "=" * 80)
    print(f"REGION:    {region}")
    print(f"SOURCE ID: {source_id}")
    print(f"FILE:      {filename}")

    if path is None:
        print("STATUS: FILE NOT FOUND")
        continue

    wb = load_workbook(
        path,
        read_only=True,
        data_only=True
    )

    print(f"SHEETS:    {wb.sheetnames}")

    for sheet_name in wb.sheetnames:

        ws = wb[sheet_name]

        headings = find_geography_headings(ws)

        print(f"\nSHEET: {sheet_name}")
        print(f"ROWS: {ws.max_row}")
        print(f"COLUMNS: {ws.max_column}")
        print(f"GEOGRAPHY BLOCKS FOUND: {len(headings)}")

        # --------------------------------------------------
        # Geography classification
        # --------------------------------------------------

        district_count = sum(
            1 for _, heading in headings
            if heading.upper().endswith(" DISTRICT")
        )

        tehsil_count = sum(
            1 for _, heading in headings
            if (
                heading.upper().endswith(" TEHSIL")
                and not heading.upper().startswith("SUB-DIVISION ")
                and not heading.upper().startswith("SUB-TEHSIL ")
        )
    )

        taluka_count = sum(
            1 for _, heading in headings
            if heading.upper().endswith(" TALUKA")
        )

        subdivision_count = sum(
            1 for _, heading in headings
            if (
                heading.upper().endswith(" SUB-DIVISION")
                or heading.upper().startswith("SUB-DIVISION ")
            )
        )

        sub_tehsil_count = sum(
            1 for _, heading in headings
            if (
                heading.upper().endswith(" SUB-TEHSIL")
                or heading.upper().startswith("SUB-TEHSIL ")
            )
        )

        protected_area_count = sum(
            1 for _, heading in headings
            if heading.upper().endswith(" PROTECTED AREA")
        )

        other_headings = [
            heading
            for _, heading in headings
            if not (
                heading.upper().endswith(" DISTRICT")
                or heading.upper().endswith(" TEHSIL")
                or heading.upper().endswith(" TALUKA")
                or heading.upper().endswith(" SUB-DIVISION")
                or heading.upper().startswith("SUB-DIVISION ")
                or heading.upper().endswith(" SUB-TEHSIL")
                or heading.upper().startswith("SUB-TEHSIL ")
                or heading.upper().endswith(" PROTECTED AREA")
            )
        ]

        print(f"DISTRICTS:       {district_count}")
        print(f"TEHSILS:         {tehsil_count}")
        print(f"TALUKAS:         {taluka_count}")
        print(f"SUB-DIVISIONS:   {subdivision_count}")
        print(f"SUB-TEHSILS:     {sub_tehsil_count}")
        print(f"PROTECTED AREAS: {protected_area_count}")
        print(f"UNKNOWN:         {len(other_headings)}")

        if other_headings:
            print("\nUNKNOWN GEOGRAPHY HEADINGS:")

            for heading in other_headings:
                print(f"  {heading}")

        # --------------------------------------------------
        # Parent-child hierarchy
        # --------------------------------------------------

        current_parent = None

        print("\nCHILD GEOGRAPHIES WITH PARENT:")

        for _, heading in headings:

            heading_upper = heading.upper()

            if heading_upper.endswith(
                (" DISTRICT", " PROTECTED AREA")
            ):
                current_parent = heading

            elif (
                heading_upper.endswith(
                    (
                        " TEHSIL",
                        " TALUKA",
                        " SUB-DIVISION",
                        " SUB-TEHSIL"
                    )
                )
                or heading_upper.startswith("SUB-DIVISION ")
                or heading_upper.startswith("SUB-TEHSIL ")
            ):
                print(f"  {current_parent} -> {heading}")

        # --------------------------------------------------
        # First geography headings
        # --------------------------------------------------

        print("\nFirst 15 geography headings:")

        for row_num, heading in headings[:15]:
            print(f"  row {row_num}: {heading}")

        # --------------------------------------------------
        # Check block lengths
        # --------------------------------------------------

        if headings:
            block_lengths = []

            for i in range(len(headings) - 1):

                current_row = headings[i][0]
                next_row = headings[i + 1][0]

                block_lengths.append(
                    next_row - current_row - 1
                )

            print("\nObservation rows between geography headings:")
            print(sorted(set(block_lengths)))

    wb.close()