from pathlib import Path

import duckdb
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

FILE_PATH = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "pbs"
    / "census_2023"
    / "table_04"
    / "table_4_punjab_districts.xlsx"
)

DATABASE_PATH = (
    PROJECT_ROOT
    / "warehouse"
    / "pakistan_census.duckdb"
)

SOURCE_SHEET = "punjab"


# These rows deliberately cover different parts of the workbook.
TEST_ROWS = [
    6,       # Attock District — ALL AGES
    99,      # Attock Tehsil — ALL AGES
    2986,    # Chiniot-area age 02 row; also had stray column-16 "s"
    12933,   # De-Excluded Area Rajanpur — ALL AGES
    16839,   # Vehari Tehsil — ALL AGES
]


def normalize_excel_value(value):
    if value is None:
        return None

    return str(value).strip()


def main():
    workbook = load_workbook(
        FILE_PATH,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook[SOURCE_SHEET]

    connection = duckdb.connect(str(DATABASE_PATH))

    failures = []

    print("RAW LOAD RECONCILIATION")
    print("=" * 70)

    for row_number in TEST_ROWS:
        excel_values = [
            normalize_excel_value(
                worksheet.cell(
                    row=row_number,
                    column=column_number,
                ).value
            )
            for column_number in range(1, 14)
        ]

        db_row = connection.execute(
            """
            select
                age_label_raw,
                all_all_sexes_raw,
                all_male_raw,
                all_female_raw,
                all_transgender_raw,
                rural_all_sexes_raw,
                rural_male_raw,
                rural_female_raw,
                rural_transgender_raw,
                urban_all_sexes_raw,
                urban_male_raw,
                urban_female_raw,
                urban_transgender_raw
            from raw.pbs_census_2023_table_04
            where source_row_number = ?
              and source_file = 'table_4_punjab_districts.xlsx'
            """,
            [row_number],
        ).fetchone()

        if db_row is None:
            failures.append(
                f"Row {row_number}: missing from DuckDB"
            )
            print(
                f"Row {row_number}: FAILED — not found in DuckDB"
            )
            continue

        db_values = list(db_row)

        if excel_values == db_values:
            print(f"Row {row_number}: PASSED")
        else:
            print(f"Row {row_number}: FAILED")

            failures.append(
                f"Row {row_number}: values do not match"
            )

            print(f"  Excel:  {excel_values}")
            print(f"  DuckDB: {db_values}")

    workbook.close()
    connection.close()

    print()
    print("=" * 70)

    if failures:
        print(f"Validation failures: {len(failures)}")

        for failure in failures:
            print(f"- {failure}")

        raise ValueError(
            "Raw reconciliation failed."
        )

    print(
        f"All {len(TEST_ROWS)} reconciliation checks passed."
    )


if __name__ == "__main__":
    main()