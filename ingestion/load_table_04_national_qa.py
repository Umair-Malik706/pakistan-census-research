from pathlib import Path
import csv

import duckdb
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

MANIFEST_PATH = PROJECT_ROOT / "metadata" / "source_manifest.csv"
RAW_DIR = PROJECT_ROOT / "data" / "raw" / "pbs" / "census_2023" / "table_04"
WAREHOUSE_PATH = PROJECT_ROOT / "warehouse" / "pakistan_census.duckdb"

SOURCE_ID = "pbs_2023_t04_national"


def find_file(filename):
    matches = list(RAW_DIR.rglob(filename))

    if not matches:
        raise FileNotFoundError(filename)

    return matches[0]


def clean_cell(value):
    if value is None:
        return None

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


with MANIFEST_PATH.open("r", encoding="utf-8-sig", newline="") as file:
    manifest = list(csv.DictReader(file))


source = next(
    row
    for row in manifest
    if row["source_id"].strip() == SOURCE_ID
)

filename = source["local_filename"].strip()
path = find_file(filename)

workbook = load_workbook(
    path,
    read_only=True,
    data_only=True
)

sheet_name = workbook.sheetnames[0]
worksheet = workbook[sheet_name]

records = []

# National workbook:
# row 5 = PAKISTAN heading
# rows 6–97 = 92 Table 4 observations
for row_number, row in enumerate(
    worksheet.iter_rows(
        min_row=6,
        max_row=97,
        min_col=1,
        max_col=13,
        values_only=True
    ),
    start=6
):

    records.append(
        (
            SOURCE_ID,
            filename,
            sheet_name,
            row_number,
            clean_cell(row[0]),
            *[clean_cell(value) for value in row[1:13]],
        )
    )


if len(records) != 92:
    raise ValueError(
        f"Expected 92 national observations, got {len(records)}"
    )


connection = duckdb.connect(str(WAREHOUSE_PATH))

connection.execute("create schema if not exists raw")

connection.execute("""
    create or replace table raw.pbs_census_2023_table_04_national_qa (
        source_id varchar,
        source_file varchar,
        source_sheet varchar,
        source_row_number integer,
        age_label_raw varchar,

        all_all_sexes_raw varchar,
        all_male_raw varchar,
        all_female_raw varchar,
        all_transgender_raw varchar,

        rural_all_sexes_raw varchar,
        rural_male_raw varchar,
        rural_female_raw varchar,
        rural_transgender_raw varchar,

        urban_all_sexes_raw varchar,
        urban_male_raw varchar,
        urban_female_raw varchar,
        urban_transgender_raw varchar
    )
""")

connection.executemany(
    """
    insert into raw.pbs_census_2023_table_04_national_qa
    values (
        ?, ?, ?, ?, ?,
        ?, ?, ?, ?,
        ?, ?, ?, ?,
        ?, ?, ?, ?
    )
    """,
    records,
)

connection.close()
workbook.close()

print(f"Loaded {len(records)} national QA observations.")