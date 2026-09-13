from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

FILE_PATH = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "pbs"
    / "census_2023"
    / "table_04"
    / "table_4_punjab_districts.xlsx"
)


def main():
    workbook = load_workbook(
        FILE_PATH,
        read_only=True,
        data_only=True,
    )

    print(f"Workbook: {FILE_PATH.name}")
    print(f"Number of sheets: {len(workbook.sheetnames)}")
    print()

    print("Sheet names:")
    for number, sheet_name in enumerate(workbook.sheetnames, start=1):
        print(f"{number}. {sheet_name}")

    print()
    print("First sheet preview:")
    print("-" * 80)

    worksheet = workbook[workbook.sheetnames[0]]

    print(f"Sheet: {worksheet.title}")
    print(f"Rows: {worksheet.max_row}")
    print(f"Columns: {worksheet.max_column}")
    print()

    for row in worksheet.iter_rows(
        min_row=1,
        max_row=min(15, worksheet.max_row),
        values_only=True,
    ):
        print(row)

    workbook.close()


if __name__ == "__main__":
    main()