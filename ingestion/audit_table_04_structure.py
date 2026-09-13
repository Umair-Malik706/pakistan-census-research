from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

FILE_PATH = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "pbs"
    / "census_2023"
    / "table_04"
    / "table_4_punjab_districts.xlsx"
)


def is_blank(value):
    return value is None or value == ""


def main():
    workbook = load_workbook(
        FILE_PATH,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook["punjab"]

    districts = []
    district_row_counts = Counter()
    age_labels = Counter()
    non_numeric_values = Counter()

    extra_column_values = Counter()

    current_district = None

    for row_number, row in enumerate(
        worksheet.iter_rows(values_only=True),
        start=1,
    ):
        values = list(row)

        first_value = values[0]

        # Check whether columns 14-16 contain any actual values.
        for column_number, value in enumerate(values[13:], start=14):
            if not is_blank(value):
                extra_column_values[column_number] += 1

        if isinstance(first_value, str):
            first_text = first_value.strip()

            # District rows appear to contain a district name
            # in column 1 and no population values.
            if (
                first_text.upper().endswith(" DISTRICT")
                and all(is_blank(value) for value in values[1:13])
            ):
                current_district = first_text
                districts.append(
                    (row_number, current_district)
                )
                continue

        # Ignore rows before the first district.
        if current_district is None:
            continue

        # Ignore completely blank rows.
        if all(is_blank(value) for value in values[:13]):
            continue

        district_row_counts[current_district] += 1

        if first_value is not None:
            age_label = str(first_value).strip()
            age_labels[age_label] += 1

        # Examine the twelve population columns.
        for value in values[1:13]:
            if is_blank(value):
                continue

            if not isinstance(value, (int, float)):
                non_numeric_values[str(value).strip()] += 1

    workbook.close()

    print("WORKBOOK STRUCTURE AUDIT")
    print("=" * 70)

    print()
    print(f"Districts detected: {len(districts)}")

    print()
    print("First 10 districts:")
    for row_number, district in districts[:10]:
        print(f"Row {row_number}: {district}")

    print()
    print("Last 5 districts:")
    for row_number, district in districts[-5:]:
        print(f"Row {row_number}: {district}")

    print()
    print("Rows per district:")
    row_count_distribution = Counter(
        district_row_counts.values()
    )

    for row_count, number_of_districts in sorted(
        row_count_distribution.items()
    ):
        print(
            f"{number_of_districts} district(s) "
            f"have {row_count} data rows"
        )

    print()
    print(f"Unique age/summary labels: {len(age_labels)}")

    print()
    print("First 30 unique labels:")
    for label in sorted(age_labels)[:30]:
        print(label)

    print()
    print("Non-numeric values found in population columns:")
    if non_numeric_values:
        for value, count in non_numeric_values.most_common():
            print(f"{repr(value)}: {count}")
    else:
        print("None")

    print()
    print("Non-empty values in columns 14-16:")
    if extra_column_values:
        for column_number, count in sorted(
            extra_column_values.items()
        ):
            print(
                f"Column {column_number}: "
                f"{count} non-empty value(s)"
            )
    else:
        print("None")


if __name__ == "__main__":
    main()