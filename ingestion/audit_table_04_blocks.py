from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

FILE_PATH = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "pbs"
    / "census_2023"
    / "table_04"
    / "table_4_punjab_districts.xlsx"
)


def is_blank(value):
    return value is None or value == ""


def main():
    workbook = load_workbook(
        FILE_PATH,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook["punjab"]

    geography_headings = []
    unusual_headings = []
    column_16_values = []

    for row_number, row in enumerate(
        worksheet.iter_rows(values_only=True),
        start=1,
    ):
        values = list(row)

        first_value = values[0]

        if isinstance(first_value, str):
            text = first_value.strip()
            upper_text = text.upper()

            if (
                text
                and all(is_blank(value) for value in values[1:13])
            ):
                if (
                    upper_text.endswith(" DISTRICT")
                    or upper_text.endswith(" TEHSIL")
                ):
                    geography_headings.append(
                        (row_number, text)
                    )

                elif row_number != 1:
                    unusual_headings.append(
                        (row_number, text)
                    )

        if len(values) >= 16 and not is_blank(values[15]):
            column_16_values.append(
                (row_number, values[15])
            )

    print("GEOGRAPHY BLOCK AUDIT")
    print("=" * 70)

    print()
    print(
        f"Geographic blocks detected: "
        f"{len(geography_headings)}"
    )

    print()
    print("Spacing between geography headings:")

    spacing_counts = {}

    for current, next_item in zip(
        geography_headings,
        geography_headings[1:],
    ):
        current_row, current_name = current
        next_row, next_name = next_item

        spacing = next_row - current_row

        spacing_counts[spacing] = (
            spacing_counts.get(spacing, 0) + 1
        )

        if spacing != 93:
            print()
            print("Unexpected spacing:")
            print(
                f"Row {current_row}: {current_name}"
            )
            print(
                f"Row {next_row}: {next_name}"
            )
            print(
                f"Difference: {spacing} rows"
            )

    print()
    print("Spacing summary:")
    for spacing, count in sorted(spacing_counts.items()):
        print(
            f"{spacing} rows apart: {count} occurrence(s)"
        )

    print()
    print("Unusual structural headings:")
    if unusual_headings:
        for row_number, text in unusual_headings:
            print(
                f"Row {row_number}: {repr(text)}"
            )
    else:
        print("None")

    print()
    print("Non-empty values in column 16:")
    if column_16_values:
        for row_number, value in column_16_values:
            print(
                f"Row {row_number}: {repr(value)}"
            )
    else:
        print("None")

    workbook.close()


if __name__ == "__main__":
    main()